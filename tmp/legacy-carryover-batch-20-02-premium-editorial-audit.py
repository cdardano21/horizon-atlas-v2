from __future__ import annotations

import hashlib
import json
import re
from collections import Counter, defaultdict
from datetime import datetime, timezone
from pathlib import Path
from urllib.parse import urlparse

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
WORKBOOK = ROOT / "data/legacy-carryover-batch-20-02/DestinationFinderAI-Legacy-Carryover-Batch-20-Premium-Enriched-Final-v3.3.xlsx"
EXPECTED_SHA256 = "015209686a8a68c20c32288498be40b88802268cc967a4ef71fbaa143b9d2a3a"
EXPECTED_KEYS = (
    "radovljica-slovenia", "osaka-japan", "sitges-spain", "estepona-spain", "lake-bled-slovenia",
    "olbia-italy", "hiroshima-japan", "kobe-japan", "alghero-italy", "hakodate-japan",
    "desenzano-del-garda-italy", "onomichi-japan", "cartagena-spain", "gijon-spain", "girona-spain",
    "ptuj-slovenia", "koper-slovenia", "murcia-spain", "takayama-japan", "dubrovnik-croatia",
)

CUSTOMER_FIELDS = {
    "DESTINATIONS": ("short_description", "long_description"),
    "NEIGHBORHOODS": ("best_for", "summary", "housing_character", "pros", "cons"),
    "PLACES": ("place_name", "description", "best_for"),
    "RESOURCES": ("description",),
    "MEDIA": ("caption", "subject"),
    "HOUSING_PROPERTY": ("restrictions_summary", "buying_process_summary", "rental_rules_notes"),
    "HEALTHCARE_INSURANCE": ("system_summary", "public_access_foreigners", "international_insurance_notes", "pharmacy_notes"),
    "VISA_RESIDENCY": ("visa_type", "residency_option", "work_rights", "renewal_notes", "permanent_residency_path", "citizenship_path"),
    "TAXES_FINANCE": ("summary", "income_tax_notes", "retirement_income_notes", "capital_gains_notes", "property_tax_notes", "vat_sales_tax_notes", "inheritance_wealth_notes", "us_tax_treaty_notes", "bank_account_foreigner_notes", "currency_notes"),
    "LGBTQ_INCLUSIVITY": ("social_acceptance", "community_scene", "pride_events", "nightlife_social", "healthcare_access", "areas_resources", "safety_considerations", "evidence_summary"),
    "SAFETY_RISKS": ("summary", "mitigation_notes"),
    "TRANSPORT_AIRPORTS": ("name", "summary", "parking_notes", "rideshare_notes"),
    "CONNECTIVITY_REMOTE_WORK": ("coworking_summary", "us_time_zone_fit", "remote_work_notes"),
    "LANGUAGE_INTEGRATION": ("integration_notes", "language_resources"),
    "PETS": ("import_requirements", "quarantine_notes", "vaccination_notes", "pet_friendly_rentals", "vet_access", "emergency_vet_access", "dog_parks_summary", "airline_notes"),
    "FAMILY_EDUCATION": ("summary", "international_schools", "childcare_notes", "universities", "pediatric_care", "family_activities"),
    "COMMUNITY_SOCIAL": ("summary", "expat_presence", "clubs_groups", "volunteering", "ease_meeting_people", "age_mix", "transient_vs_rooted"),
    "ACCESSIBILITY": ("wheelchair_access", "sidewalk_quality", "hills_terrain", "accessible_transit", "elevator_access", "medical_equipment", "mobility_notes"),
    "BUREAUCRACY_SETUP": ("summary", "typical_documents", "estimated_timeline"),
    "WORK_BUSINESS": ("major_industries", "employment_notes", "work_authorization", "entrepreneurship", "business_formation", "coworking", "remote_work_suitability"),
    "RETIREMENT_AGING": ("medicare_notes", "social_security_notes", "senior_discounts", "assisted_living", "home_healthcare", "aging_in_place", "retirement_notes"),
    "LIFESTYLE_LAWS": ("summary", "important_rules"),
    "REALITY_CHECK": ("title", "detail"),
    "MOVE_CHECKLIST": ("task", "description"),
    "ENVIRONMENT_QUALITY": ("air_quality_summary", "water_quality_summary", "heat_humidity_comfort", "noise_summary", "light_pollution_summary", "mosquito_pest_pressure", "wildfire_smoke_exposure", "drought_water_stress", "environmental_notes"),
    "DAILY_LIFE_PRACTICALITY": ("car_need", "driving_difficulty", "parking_difficulty", "grocery_access", "pharmacy_access", "fitness_wellness_access", "banking_practicality", "card_payment_acceptance", "cash_usage", "mobile_payment_usage", "delivery_services", "emergency_services_summary", "senior_services_summary", "childcare_access", "newcomer_friction", "things_residents_wish_they_knew"),
    "EVENTS_SEASONALITY": ("name", "description", "weather_context", "best_for", "avoid_if"),
    "LIFESTYLE_FEATURES": ("display_name", "evidence_summary"),
}

DESTINATION_SPECIFIC_MODULES = {
    "DESTINATIONS": ("short_description", "long_description"),
    "REALITY_CHECK": ("title", "detail"),
    "DAILY_LIFE_PRACTICALITY": CUSTOMER_FIELDS["DAILY_LIFE_PRACTICALITY"],
    "RETIREMENT_AGING": CUSTOMER_FIELDS["RETIREMENT_AGING"],
    "FAMILY_EDUCATION": CUSTOMER_FIELDS["FAMILY_EDUCATION"],
    "COMMUNITY_SOCIAL": CUSTOMER_FIELDS["COMMUNITY_SOCIAL"],
    "CONNECTIVITY_REMOTE_WORK": CUSTOMER_FIELDS["CONNECTIVITY_REMOTE_WORK"],
    "LANGUAGE_INTEGRATION": CUSTOMER_FIELDS["LANGUAGE_INTEGRATION"],
    "ACCESSIBILITY": CUSTOMER_FIELDS["ACCESSIBILITY"],
    "TRANSPORT_AIRPORTS": CUSTOMER_FIELDS["TRANSPORT_AIRPORTS"],
    "WORK_BUSINESS": CUSTOMER_FIELDS["WORK_BUSINESS"],
}

PROHIBITED_PATTERNS = {
    "concrete reference": re.compile(r"\bconcrete reference\b", re.IGNORECASE),
    "practical area to compare": re.compile(r"\bpractical area to compare\b", re.IGNORECASE),
    "KYC": re.compile(r"\bKYC\b", re.IGNORECASE),
    "regional universities vary": re.compile(r"\bregional universities vary\b", re.IGNORECASE),
    "test destination for medical access": re.compile(r"\btest\b.{0,100}\bfor medical access\b", re.IGNORECASE),
    "binary legal token": re.compile(r"\bbinary legal token\b", re.IGNORECASE),
    "Provider-specific": re.compile(r"\bprovider-specific\b", re.IGNORECASE),
    "Permit-specific": re.compile(r"\bpermit-specific\b", re.IGNORECASE),
    "implementation jargon": re.compile(r"\b(?:parser|adapter|runtime|database|workbook)\b", re.IGNORECASE),
}


def text(value: object) -> str:
    return "" if value is None else str(value).strip()


def records(workbook, sheet_name: str) -> list[dict[str, object]]:
    rows = workbook[sheet_name].iter_rows(values_only=True)
    headers = [text(value) for value in next(rows, ())]
    return [dict(zip(headers, row)) for row in rows if any(value is not None for value in row)]


def host(value: object) -> str:
    return urlparse(text(value)).netloc.lower().removeprefix("www.")


def valid_url(value: object) -> bool:
    parsed = urlparse(text(value))
    return parsed.scheme in {"http", "https"} and bool(parsed.netloc)


def truthy(value: object) -> bool:
    return text(value).lower() in {"true", "1", "yes"}


def main() -> None:
    errors: list[str] = []
    notes: list[str] = []
    workbook_sha256 = hashlib.sha256(WORKBOOK.read_bytes()).hexdigest()
    if workbook_sha256 != EXPECTED_SHA256:
        errors.append(f"WORKBOOK_SHA256_MISMATCH:{workbook_sha256}")

    workbook = load_workbook(WORKBOOK, read_only=True, data_only=True)
    sheet_rows = {name: records(workbook, name) for name in workbook.sheetnames}
    destinations = sheet_rows["DESTINATIONS"]
    destination_keys = [text(row.get("destination_key")) for row in destinations]
    if tuple(destination_keys) != EXPECTED_KEYS:
        errors.append("DESTINATION_SCOPE_OR_ORDER_MISMATCH")

    expected_counts = {
        "DESTINATIONS": 20,
        "NEIGHBORHOODS": 100,
        "PLACES": 360,
        "LIFESTYLE_FEATURES": 340,
        "REALITY_CHECK": 120,
        "DAILY_LIFE_PRACTICALITY": 20,
        "RETIREMENT_AGING": 20,
        "FAMILY_EDUCATION": 40,
        "COMMUNITY_SOCIAL": 60,
        "CONNECTIVITY_REMOTE_WORK": 20,
        "LANGUAGE_INTEGRATION": 20,
        "ACCESSIBILITY": 20,
        "TRANSPORT_AIRPORTS": 60,
        "WORK_BUSINESS": 20,
    }
    for sheet_name, expected_count in expected_counts.items():
        if len(sheet_rows[sheet_name]) != expected_count:
            errors.append(f"ROW_COUNT:{sheet_name}:{len(sheet_rows[sheet_name])}:{expected_count}")

    overview_profiles = sum(
        bool(text(row.get("short_description")) and text(row.get("long_description")))
        for row in destinations
    )
    if overview_profiles != 20:
        errors.append(f"OVERVIEW_PROFILE_COVERAGE:{overview_profiles}:20")

    neighborhoods = sheet_rows["NEIGHBORHOODS"]
    neighborhood_descriptions = sum(
        bool(text(row.get("neighborhood_name")) and text(row.get("summary")))
        for row in neighborhoods
    )
    if neighborhood_descriptions != 100:
        errors.append(f"NEIGHBORHOOD_DESCRIPTION_COVERAGE:{neighborhood_descriptions}:100")
    neighborhood_keys = {
        (text(row.get("destination_key")), text(row.get("neighborhood_key")))
        for row in neighborhoods
    }

    destination_tourism_hosts = {
        text(row.get("destination_key")): host(row.get("official_tourism_url"))
        for row in destinations
    }
    places = sheet_rows["PLACES"]
    named_places = sum(bool(text(row.get("place_name")) and text(row.get("description"))) for row in places)
    website_coverage = sum(valid_url(row.get("website_url")) for row in places)
    neighborhood_coverage = sum(
        (text(row.get("destination_key")), text(row.get("neighborhood_key"))) in neighborhood_keys
        for row in places
    )
    maps_coverage = sum(valid_url(row.get("google_maps_url")) for row in places)
    google_primary_urls = sum("google." in host(row.get("website_url")) or "goo.gl" in host(row.get("website_url")) for row in places)
    verified_primary_urls = sum(
        truthy(row.get("verified"))
        and text(row.get("confidence")).upper() == "HIGH"
        and host(row.get("website_url")) != destination_tourism_hosts.get(text(row.get("destination_key")))
        for row in places
    )
    tourism_fallback_urls = sum(
        not truthy(row.get("verified"))
        and text(row.get("confidence")).upper() == "MEDIUM"
        and host(row.get("website_url")) == destination_tourism_hosts.get(text(row.get("destination_key")))
        for row in places
    )
    if named_places != 360:
        errors.append(f"NAMED_PLACE_COVERAGE:{named_places}:360")
    if website_coverage != 360 or google_primary_urls:
        errors.append(f"PRIMARY_WEBSITE_POLICY:{website_coverage}:360:GOOGLE_PRIMARY={google_primary_urls}")
    if neighborhood_coverage != 360:
        errors.append(f"PLACE_NEIGHBORHOOD_COVERAGE:{neighborhood_coverage}:360")
    if maps_coverage != 360:
        errors.append(f"SECONDARY_MAPS_COVERAGE:{maps_coverage}:360")
    if verified_primary_urls + tourism_fallback_urls != 360:
        errors.append(f"WEBSITE_VERIFICATION_OR_FALLBACK_POLICY:{verified_primary_urls}+{tourism_fallback_urls}:360")

    module_specificity: dict[str, dict[str, int]] = {}
    for sheet_name, fields in DESTINATION_SPECIFIC_MODULES.items():
        by_destination: dict[str, list[str]] = defaultdict(list)
        for row in sheet_rows[sheet_name]:
            destination_key = text(row.get("destination_key"))
            values = [text(row.get(field)) for field in fields if text(row.get(field))]
            if values:
                by_destination[destination_key].append(" | ".join(values))
        signatures = {
            destination_key: hashlib.sha256("\n".join(values).encode()).hexdigest()
            for destination_key, values in by_destination.items()
        }
        module_specificity[sheet_name] = {
            "destinationsCovered": len(signatures),
            "uniqueDestinationCopySignatures": len(set(signatures.values())),
        }
        if len(signatures) != 20 or len(set(signatures.values())) != 20:
            errors.append(f"DESTINATION_SPECIFIC_COPY:{sheet_name}:{len(signatures)}:{len(set(signatures.values()))}")

    prohibited_matches: list[dict[str, object]] = []
    customer_cells_scanned = 0
    for sheet_name, fields in CUSTOMER_FIELDS.items():
        for row_number, row in enumerate(sheet_rows[sheet_name], start=2):
            for field in fields:
                value = text(row.get(field))
                if not value:
                    continue
                customer_cells_scanned += 1
                for label, pattern in PROHIBITED_PATTERNS.items():
                    if pattern.search(value):
                        prohibited_matches.append({
                            "pattern": label,
                            "sheet": sheet_name,
                            "field": field,
                            "row": row_number,
                            "destinationKey": text(row.get("destination_key")),
                            "value": value,
                        })
    if prohibited_matches:
        errors.append(f"PROHIBITED_CUSTOMER_COPY:{len(prohibited_matches)}")

    media = sheet_rows["MEDIA"]
    media_verified = sum(truthy(row.get("verified")) for row in media)
    media_with_license_notes = sum(bool(text(row.get("license_notes"))) for row in media)
    if media_verified != len(media):
        notes.append("Media rows remain subject to file-level rights verification before public display.")

    result = {
        "status": "PASS" if not errors else "FAIL",
        "readOnly": True,
        "createdAt": datetime.now(timezone.utc).isoformat(),
        "workbook": {"path": str(WORKBOOK.relative_to(ROOT)), "sha256": workbook_sha256},
        "errors": errors,
        "notes": notes,
        "coverage": {
            "overviewProfiles": overview_profiles,
            "neighborhoodRows": len(neighborhoods),
            "neighborhoodDescriptions": neighborhood_descriptions,
            "namedRecommendationRows": named_places,
            "recommendationWebsiteUrls": website_coverage,
            "recommendationNeighborhoodAssociations": neighborhood_coverage,
            "secondaryGoogleMapsUrls": maps_coverage,
            "verifiedVenuePrimaryUrls": verified_primary_urls,
            "officialDestinationFallbackUrls": tourism_fallback_urls,
            "lifestyleRows": len(sheet_rows["LIFESTYLE_FEATURES"]),
            "realityCheckRows": len(sheet_rows["REALITY_CHECK"]),
        },
        "moduleSpecificity": module_specificity,
        "customerCopy": {
            "cellsScanned": customer_cells_scanned,
            "prohibitedMatches": prohibited_matches,
        },
        "media": {
            "rows": len(media),
            "verifiedRows": media_verified,
            "unverifiedRows": len(media) - media_verified,
            "rowsWithLicenseNotes": media_with_license_notes,
            "publicationReady": media_verified == len(media),
        },
        "placeConfidence": dict(sorted(Counter(text(row.get("confidence")) for row in places).items())),
    }
    print(json.dumps(result, indent=2, ensure_ascii=True))
    if errors:
        raise SystemExit(1)


if __name__ == "__main__":
    main()