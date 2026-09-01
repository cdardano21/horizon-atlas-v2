import openpyxl

WORKBOOK_PATH = "data/legacy-migration-batch-20/DestinationFinderAI_Legacy_Production_Batch_20_Reader_Ready_Authoritative_v3.3.xlsx"

KEYS = [
    "ajijic-mexico", "boquete-panama", "chiang-mai-thailand", "cuenca-ecuador", "da-nang-vietnam",
    "florianopolis-brazil", "funchal-portugal", "george-town-malaysia", "hua-hin-thailand", "lucca-italy",
    "merida-mexico", "monopoli-italy", "montevideo-uruguay", "nafplio-greece", "nice-france",
    "palm-springs-california-united-states", "paphos-cyprus", "santander-spain",
    "savannah-georgia-united-states", "sibenik-croatia",
]

PHRASES = [
    "before publication",
    "must be extracted",
    "must be rechecked",
    "must be checked before",
    "must be refreshed before",
    "should be confirmed directly before",
    "require department-level confirmation",
    "research seed",
    "specific items to confirm",
]

wb = openpyxl.load_workbook(WORKBOOK_PATH, data_only=True)
hits = 0
for sheet_name in wb.sheetnames:
    ws = wb[sheet_name]
    header = [c.value for c in ws[1]] if ws.max_row >= 1 else []
    if "destination_key" not in header:
        continue
    idx = {h: i for i, h in enumerate(header) if h}
    dk_col = idx["destination_key"]
    for row_num, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if not row or dk_col >= len(row):
            continue
        dk = row[dk_col]
        if dk not in KEYS:
            continue
        for col_name, col_i in idx.items():
            if col_i >= len(row):
                continue
            value = row[col_i]
            if isinstance(value, str):
                lower = value.lower()
                for phrase in PHRASES:
                    if phrase in lower:
                        hits += 1
                        print(f"[{sheet_name}] row={row_num} dest={dk} col={col_name} phrase={phrase!r}")
                        print(f"    VALUE: {value!r}")
                        print()
print(f"Total hits: {hits}")
