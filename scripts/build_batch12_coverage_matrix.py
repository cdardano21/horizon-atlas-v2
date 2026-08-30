#!/usr/bin/env python3
"""
Builds a machine-readable coverage matrix from the two final, committed Batch 1 / Batch 2
v3.3 workbooks. Read-only - never modifies either workbook. This is the minimum-completeness
benchmark for legacy-migration-pilot-06 Phase 5D (full six-destination population).

For every sheet, records: headers, row count, and per-destination nonblank-field counts (for
destination-scoped sheets that have a destination_key column). For LIFESTYLE_FEATURES, also
records the controlled feature_key taxonomy actually in use across both batches.

Usage:
  python3 scripts/build_batch12_coverage_matrix.py
"""
import json
from pathlib import Path

import openpyxl

REPO_ROOT = Path(__file__).resolve().parent.parent
BATCH_WORKBOOKS = {
    "batch01_v3.3": REPO_ROOT / "data" / "DestinationFinderAI_Expansion_Batch_01_5_Destinations_v3.3.xlsx",
    "batch02_v3.3": REPO_ROOT / "data" / "DestinationFinderAI_Expansion_Batch_02_5_Destinations_v3.3.xlsx",
}
OUTPUT_PATH = REPO_ROOT / "data" / "legacy-migration-pilot-06" / "batch12-coverage-matrix.json"


def nonblank(v):
    return v is not None and str(v).strip() != ""


def analyze_workbook(path):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    sheets = {}
    for name in wb.sheetnames:
        ws = wb[name]
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            sheets[name] = {"headers": [], "rowCount": 0}
            continue
        header = [h for h in rows[0] if h is not None]
        # openpyxl's read_only used-range can overshoot true data (trailing formatted-but-empty
        # rows) - filter to genuinely nonblank rows only, matching the validator's own convention.
        data_rows = [r for r in rows[1:] if any(nonblank(v) for v in r)]
        entry = {"headers": header, "rowCount": len(data_rows)}
        if "destination_key" in header:
            dk_idx = header.index("destination_key")
            per_dest_nonblank = {}
            per_dest_rowcount = {}
            for r in data_rows:
                dk = r[dk_idx] if dk_idx < len(r) else None
                if not nonblank(dk):
                    continue
                per_dest_rowcount[dk] = per_dest_rowcount.get(dk, 0) + 1
                nb = sum(1 for i, v in enumerate(r) if i < len(header) and nonblank(v))
                per_dest_nonblank.setdefault(dk, []).append(nb)
            entry["rowCountByDestination"] = per_dest_rowcount
            entry["avgNonblankFieldsByDestination"] = {
                k: round(sum(v) / len(v), 1) for k, v in per_dest_nonblank.items()
            }
        if name == "LIFESTYLE_FEATURES" and "feature_key" in header:
            fk_idx = header.index("feature_key")
            feature_keys = sorted({r[fk_idx] for r in data_rows if fk_idx < len(r) and nonblank(r[fk_idx])})
            entry["controlledFeatureKeys"] = feature_keys
        if name == "DESTINATIONS" and "destination_key" in header:
            dk_idx = header.index("destination_key")
            entry["destinationKeys"] = sorted({r[dk_idx] for r in data_rows if dk_idx < len(r) and nonblank(r[dk_idx])})
        sheets[name] = entry
    wb.close()
    return {"sheetOrder": wb.sheetnames, "sheets": sheets}


def main():
    matrix = {}
    for label, path in BATCH_WORKBOOKS.items():
        if not path.exists():
            raise SystemExit(f"Missing workbook: {path}")
        matrix[label] = analyze_workbook(path)

    OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    OUTPUT_PATH.write_text(json.dumps(matrix, indent=2, default=str), encoding="utf-8")
    print(f"Wrote {OUTPUT_PATH.relative_to(REPO_ROOT)}")
    for label, data in matrix.items():
        print(f"\n{label}: {len(data['sheetOrder'])} sheets")


if __name__ == "__main__":
    main()
