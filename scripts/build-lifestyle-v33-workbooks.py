"""
Permanent, reusable builder for the Lifestyle v3.3 candidate workbooks.

Reads final evidence exclusively from the tracked canonical directory
data/lifestyle-v33/ (never from tmp/ scratch work). For each batch:
  1. Load the v3.2 source read-only; save to a NEW v3.3 path (the source is
     never opened for writing and never saved back to its own path).
  2. Append LIFESTYLE_FEATURES as the final sheet, with the exact approved
     16-column header, populated with that batch's final validated evidence
     rows (data/lifestyle-v33/batchNN-lifestyle-evidence.json).
  3. Append a small number of already-reviewed, approved new PLACES rows
     (additive only - inserted into the first genuinely blank rows right
     after the existing data, never touching any existing row).
  4. Bump WORKBOOK_METADATA.schema_version "3.2" -> "3.3" (the only change to
     any pre-existing row in either workbook).

Every other existing sheet/row/column/value/hyperlink is left exactly as
openpyxl loaded it. DESTINATION_SCORES is not touched.

Run from repo root: python3 scripts/build-lifestyle-v33-workbooks.py
"""
import json
import os
import openpyxl

REPO_ROOT = os.path.join(os.path.dirname(__file__), "..")
EVIDENCE_DIR = os.path.join(REPO_ROOT, "data", "lifestyle-v33")

LIFESTYLE_FEATURES_HEADERS = [
    "record_key",
    "destination_key",
    "feature_group",
    "feature_key",
    "feature_value",
    "availability_level",
    "proximity_band",
    "display_label",
    "evidence_summary",
    "source_name",
    "source_url",
    "source_as_of_date",
    "confidence",
    "matching_enabled",
    "display_enabled",
    "notes",
]

PLACES_HEADER = [
    "place_key", "destination_key", "neighborhood_key", "category_key", "place_name",
    "subcategory", "description", "address", "latitude", "longitude", "price_level",
    "website_url", "google_maps_url", "phone", "best_for", "display_order",
    "source_name", "source_url", "verified", "verified_at", "confidence",
]

VERIFIED_AT = "2026-08-27"

# ---------------------------------------------------------------------------
# Approved new PLACES rows only (reviewed: MEDIUM/HIGH confidence, adequate
# evidence, not already present, valid existing category, valid existing
# neighborhood, Website/Maps fields follow each workbook's own convention).
# ---------------------------------------------------------------------------

BATCH01_NEW_PLACES = [
    {
        "place_key": "puerto-vallarta-mx-playa-olas-altas",
        "destination_key": "puerto-vallarta-mx",
        "neighborhood_key": "puerto-vallarta-mx-zona-romantica",
        "category_key": "beach",
        "place_name": "Playa Olas Altas",
        "subcategory": "Beach",
        "description": "Popular swimming beach adjoining the Zona Romantica/Los Muertos beachfront.",
        "address": None, "latitude": None, "longitude": None, "price_level": None,
        "website_url": None,
        "google_maps_url": "https://www.google.com/maps/search/?api=1&query=Playa+Olas+Altas%2C+Puerto+Vallarta",
        "phone": None,
        "best_for": "Residents and visitors seeking a verified local option",
        "display_order": 26,
        "source_name": "Wikipedia",
        "source_url": "https://en.wikipedia.org/wiki/Puerto_Vallarta",
        "verified": True,
        "verified_at": VERIFIED_AT,
        "confidence": "high",
    },
    {
        "place_key": "hoi-an-vn-cham-islands-cu-lao-cham",
        "destination_key": "hoi-an-vn",
        "neighborhood_key": "hoi-an-vn-cua-dai",
        "category_key": "water_recreation",
        "place_name": "Cham Islands (Cu Lao Cham)",
        "subcategory": "Water Recreation",
        "description": (
            "Offshore UNESCO Biosphere Reserve archipelago (~15-19km from Hoi An, reached by "
            "an approximately 30-minute speedboat crossing typically departing from Cua Dai "
            "Beach), with facilities for camping, swimming and scuba diving."
        ),
        "address": None, "latitude": None, "longitude": None, "price_level": None,
        "website_url": None,
        "google_maps_url": "https://www.google.com/maps/search/?api=1&query=Cham+Islands%2C+Hoi+An",
        "phone": None,
        "best_for": "Residents and visitors seeking a verified local option",
        "display_order": 26,
        "source_name": "Wikipedia",
        "source_url": "https://en.wikipedia.org/wiki/Ch%C3%A0m_Islands",
        "verified": True,
        "verified_at": VERIFIED_AT,
        "confidence": "high",
    },
]

BATCH02_NEW_PLACES = [
    {
        "place_key": "ascoli-piceno-it-place-21",
        "destination_key": "ascoli-piceno-it",
        "neighborhood_key": "ascoli-piceno-it-porta-romana-porta-cartara",
        "category_key": "water_recreation",
        "place_name": "River Castellano swimming area (Ponte di Cecco vicinity)",
        "subcategory": "water_recreation",
        "description": "Traditional summer swimming/bathing stretch of the River Castellano near the historic center, close to Ponte di Cecco.",
        "address": None, "latitude": None, "longitude": None, "price_level": None,
        "website_url": None,
        "google_maps_url": "https://www.google.com/maps/search/?api=1&query=Fiume+Castellano+Ascoli+Piceno+Italy",
        "phone": None,
        "best_for": "Destination recommendation",
        "display_order": 21,
        "source_name": "Wikipedia",
        "source_url": "https://en.wikipedia.org/wiki/Ascoli_Piceno",
        "verified": True,
        "verified_at": VERIFIED_AT,
        "confidence": "MEDIUM",
    },
]

BATCHES = [
    {
        "label": "Batch #1",
        "src": os.path.join(REPO_ROOT, "data/DestinationFinderAI_Expansion_Batch_01_5_Destinations_v3.2.xlsx"),
        "dst": os.path.join(REPO_ROOT, "data/DestinationFinderAI_Expansion_Batch_01_5_Destinations_v3.3.xlsx"),
        "evidence_json": os.path.join(EVIDENCE_DIR, "batch01-lifestyle-evidence.json"),
        "new_places": BATCH01_NEW_PLACES,
    },
    {
        "label": "Batch #2",
        "src": os.path.join(REPO_ROOT, "data/DestinationFinderAI_Expansion_Batch_02_5_Destinations_v3.2.xlsx"),
        "dst": os.path.join(REPO_ROOT, "data/DestinationFinderAI_Expansion_Batch_02_5_Destinations_v3.3.xlsx"),
        "evidence_json": os.path.join(EVIDENCE_DIR, "batch02-lifestyle-evidence.json"),
        "new_places": BATCH02_NEW_PLACES,
    },
]


def set_schema_version(ws, version: str) -> bool:
    for row in ws.iter_rows(min_row=2):
        key_cell = row[0]
        if key_cell.value == "schema_version":
            row[1].value = version
            return True
    return False


def find_places_next_free_row(ws) -> int:
    last_data_row = 1
    for r in range(1, ws.max_row + 1):
        if ws.cell(row=r, column=1).value:
            last_data_row = r
    return last_data_row + 1


def build(cfg: dict) -> dict:
    wb = openpyxl.load_workbook(cfg["src"])
    sheets_before = list(wb.sheetnames)

    with open(cfg["evidence_json"], "r", encoding="utf-8") as fh:
        evidence_rows = json.load(fh)

    lifestyle_ws = wb.create_sheet("LIFESTYLE_FEATURES")
    lifestyle_ws.append(LIFESTYLE_FEATURES_HEADERS)
    for row in evidence_rows:
        lifestyle_ws.append([row.get(col) for col in LIFESTYLE_FEATURES_HEADERS])

    places_ws = wb["PLACES"]
    next_row = find_places_next_free_row(places_ws)
    inserted_place_keys = []
    for i, place in enumerate(cfg["new_places"]):
        r = next_row + i
        for c, col_name in enumerate(PLACES_HEADER, start=1):
            places_ws.cell(row=r, column=c).value = place.get(col_name)
        inserted_place_keys.append(place["place_key"])

    metadata_ws = wb["WORKBOOK_METADATA"]
    schema_version_updated = set_schema_version(metadata_ws, "3.3")

    wb.save(cfg["dst"])

    return {
        "label": cfg["label"],
        "src": cfg["src"],
        "dst": cfg["dst"],
        "sheets_before": sheets_before,
        "sheets_after": list(wb.sheetnames),
        "lifestyle_rows_inserted": len(evidence_rows),
        "places_rows_inserted": len(cfg["new_places"]),
        "inserted_place_keys": inserted_place_keys,
        "places_first_new_row": next_row,
        "schema_version_updated": schema_version_updated,
    }


if __name__ == "__main__":
    results = [build(cfg) for cfg in BATCHES]
    for result in results:
        print(f"=== {result['label']}: {result['dst']} ===")
        print(f"  sheets before: {len(result['sheets_before'])}  sheets after: {len(result['sheets_after'])}")
        print(f"  new sheet(s): {[s for s in result['sheets_after'] if s not in result['sheets_before']]}")
        print(f"  LIFESTYLE_FEATURES rows inserted: {result['lifestyle_rows_inserted']}")
        print(f"  PLACES rows inserted: {result['places_rows_inserted']} starting at row {result['places_first_new_row']}: {result['inserted_place_keys']}")
        print(f"  schema_version updated: {result['schema_version_updated']}")
        print()
