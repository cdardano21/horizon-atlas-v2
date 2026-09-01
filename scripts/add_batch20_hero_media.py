"""
One-off, contained enrichment script: registers exactly one verified, license-safe hero image per
Batch 20 destination into the MEDIA sheet, which was entirely empty (zero rows, for any
destination, in the whole workbook) prior to this run. Every image is a real Wikimedia Commons
photograph of the destination itself, selected and license-checked via the Commons API before this
script was written - see the accompanying session notes for the exact search/verification steps.

Never touches any other sheet, any other destination_key, or any existing MEDIA row (there were
none). Writes only: media_key, destination_key, media_type, image_url, caption, subject,
primary_image, gallery_order, license_notes, source_name, source_url, verified, verified_at,
confidence.
"""

import openpyxl
from datetime import date

WORKBOOK_PATH = "data/legacy-migration-batch-20/DestinationFinderAI_Legacy_Production_Batch_20_Reader_Ready_Authoritative_v3.3.xlsx"

# (destination_key, destination_name, Commons file title without "File:" prefix, license short name)
HEROES = [
    ("ajijic-mexico", "Ajijic", "Ajijic Malecon Sunset.jpg", "CC BY-SA 4.0"),
    ("boquete-panama", "Boquete", "Aerial view of Boquete, Panama.jpg", "CC BY-SA 4.0"),
    ("chiang-mai-thailand", "Chiang Mai", "Chiang Mai - East gate of the city wall - 0001.jpg", "CC BY-SA 3.0"),
    ("cuenca-ecuador", "Cuenca", "Cuenca Calle Juan Jaramillo 6-01.jpg", "CC BY-SA 4.0"),
    ("da-nang-vietnam", "Da Nang", "Da-Nang Vietnam Fisher-boats-01.jpg", "CC BY-SA 3.0"),
    ("florianopolis-brazil", "Florianópolis", "Praia Mole em Florianopolis.jpg", "CC0"),
    ("funchal-portugal", "Funchal", "Funchal (Madeira, Portugal), Rua de Santa Maria 39 -- 2025 -- 0831.jpg", "CC BY-SA 4.0"),
    ("george-town-malaysia", "George Town", "Penang Road, George Town, Penang.jpg", "CC BY-SA 4.0"),
    ("hua-hin-thailand", "Hua Hin", "Hua Hin Sunrise.jpg", "CC0"),
    ("lucca-italy", "Lucca", "Walls of Lucca, May 2013 (02).JPG", "CC BY-SA 3.0"),
    ("merida-mexico", "Mérida", "Centro Cultural Teatro Yucatan - Merida, Yucatan, Mexico - Abril 2021.jpg", "CC BY 2.0"),
    ("monopoli-italy", "Monopoli", "Cala Diavolo, Monopoli Puglia (Italia).jpg", "CC BY-SA 4.0"),
    ("montevideo-uruguay", "Montevideo", "Montevideo, Uruguay (3483934133).jpg", "CC BY 2.0"),
    ("nafplio-greece", "Nafplio", "Nafplio, Greece 2022.jpg", "CC BY-SA 4.0"),
    ("nice-france", "Nice", "Promenade des Anglais (Nice), France.jpg", "CC BY-SA 4.0"),
    ("palm-springs-california-united-states", "Palm Springs", "Downtown Palm Springs CA.JPG", "CC BY-SA 3.0"),
    ("paphos-cyprus", "Paphos", "Boats near Paphos, Cyprus 1626p.jpg", "CC BY-SA 4.0"),
    ("santander-spain", "Santander", "Cantabria. Santander. Spain (3380000974).jpg", "CC BY-SA 2.0"),
    ("savannah-georgia-united-states", "Savannah", "Savannah Historic District (Savannah, Georgia) 3 10.JPG", "CC BY-SA 4.0"),
    ("sibenik-croatia", "Šibenik", "Sibenik katedrala04 Croatia.jpg", "Public domain"),
]

def file_path_url(filename: str) -> str:
    from urllib.parse import quote
    return f"https://commons.wikimedia.org/wiki/Special:FilePath/{quote(filename)}"

def file_page_url(filename: str) -> str:
    from urllib.parse import quote
    return f"https://commons.wikimedia.org/wiki/File:{quote(filename)}"

def main():
    wb = openpyxl.load_workbook(WORKBOOK_PATH)
    ws = wb["MEDIA"]
    header = [cell.value for cell in ws[1]]
    idx = {h: i for i, h in enumerate(header) if h}

    # Confirm every row in the sheet is genuinely blank (destination_key empty) before writing -
    # refuse to run if that assumption is ever violated by a future workbook change.
    for row in ws.iter_rows(min_row=2, values_only=False):
        if row[idx["destination_key"]].value:
            raise SystemExit(f"Row {row[0].row} already has a destination_key - aborting to avoid overwriting real data.")

    today = date.today().isoformat()
    for row_offset, (dest_key, dest_name, filename, license_name) in enumerate(HEROES):
        row_number = 2 + row_offset
        row = ws[row_number]
        row[idx["media_key"]].value = f"{dest_key}-hero-01"
        row[idx["destination_key"]].value = dest_key
        row[idx["media_type"]].value = "photo"
        row[idx["image_url"]].value = file_path_url(filename)
        row[idx["caption"]].value = f"{dest_name}"
        row[idx["subject"]].value = dest_name
        row[idx["primary_image"]].value = "YES"
        row[idx["gallery_order"]].value = 1
        row[idx["license_notes"]].value = f"{license_name}, via Wikimedia Commons"
        row[idx["source_name"]].value = "Wikimedia Commons"
        row[idx["source_url"]].value = file_page_url(filename)
        row[idx["verified"]].value = True
        row[idx["verified_at"]].value = today
        if "confidence" in idx:
            row[idx["confidence"]].value = "HIGH"

    wb.save(WORKBOOK_PATH)
    print(f"Wrote {len(HEROES)} hero MEDIA rows to {WORKBOOK_PATH}")

if __name__ == "__main__":
    main()
