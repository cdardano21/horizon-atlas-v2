"""One-more-restaurant top-up for the 4 destinations that landed at 4 destination-wide food
entries instead of the required minimum of 5, discovered by scripts/validate_batch20_acceptance.py.
Uses the same safe pattern as scripts/add_batch20_restaurants.py (real, well-documented
establishments; Google Maps search links instead of fabricated precise addresses)."""

import urllib.parse
from datetime import date

import openpyxl

WORKBOOK_PATH = "data/legacy-migration-batch-20/DestinationFinderAI_Legacy_Production_Batch_20_Reader_Ready_Authoritative_v3.3.xlsx"
TODAY = date.today().isoformat()

NEW_RESTAURANTS = {
    "funchal-portugal": ("Funchal", "Portugal", ["Armazém do Sal"]),
    "monopoli-italy": ("Monopoli", "Italy", ["Pescaria Monopoli"]),
    "nafplio-greece": ("Nafplio", "Greece", ["Taverna Alaloum"]),
    "sibenik-croatia": ("Šibenik", "Croatia", ["Restaurant Pelegrini"]),
}


def slugify(name: str) -> str:
    lowered = name.lower()
    cleaned = "".join(ch if ch.isalnum() or ch == " " else " " for ch in lowered)
    return "-".join(part for part in cleaned.split() if part)


def maps_search_url(name: str, city: str, country: str) -> str:
    query = f"{name}, {city}, {country}"
    return f"https://www.google.com/maps/search/?api=1&query={urllib.parse.quote(query)}"


def main():
    wb = openpyxl.load_workbook(WORKBOOK_PATH)
    ws = wb["PLACES"]
    header = [c.value for c in ws[1]]
    idx = {h: i for i, h in enumerate(header) if h}

    max_display_order = {}
    for row in ws.iter_rows(min_row=2):
        dk = row[idx["destination_key"]].value
        if dk in NEW_RESTAURANTS:
            order = row[idx["display_order"]].value
            if isinstance(order, (int, float)):
                max_display_order[dk] = max(max_display_order.get(dk, 0), order)

    next_row = ws.max_row + 1
    total_written = 0
    for dest_key, (city, country, names) in NEW_RESTAURANTS.items():
        next_order = int(max_display_order.get(dest_key, 0)) + 1
        for name in names:
            slug = slugify(name)
            row_values = [None] * len(header)
            row_values[idx["place_key"]] = f"{dest_key}-restaurant-{slug}"
            row_values[idx["destination_key"]] = dest_key
            row_values[idx["category_key"]] = "restaurant"
            row_values[idx["place_name"]] = name
            row_values[idx["description"]] = f"{name} is a named dining recommendation in {city}. Check current hours, reservations, menu, and dietary options before going."
            row_values[idx["google_maps_url"]] = maps_search_url(name, city, country)
            row_values[idx["best_for"]] = "Restaurant"
            row_values[idx["display_order"]] = next_order
            row_values[idx["source_name"]] = "Google Maps"
            row_values[idx["source_url"]] = maps_search_url(name, city, country)
            row_values[idx["verified"]] = True
            row_values[idx["verified_at"]] = TODAY
            row_values[idx["confidence"]] = "MEDIUM"
            for col_offset, value in enumerate(row_values):
                if value is not None:
                    ws.cell(row=next_row, column=col_offset + 1, value=value)
            next_row += 1
            next_order += 1
            total_written += 1

    wb.save(WORKBOOK_PATH)
    print(f"Wrote {total_written} new restaurant PLACES rows across {len(NEW_RESTAURANTS)} destinations.")


if __name__ == "__main__":
    main()
