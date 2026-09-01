"""
Bounded Batch 20 workbook cleanup: blanks reader-facing text cells across the whole workbook that
contain internal research/production-status language (the exact phrases the task calls out), for
the 20 Batch 20 destinations only. Never touches any other destination_key.

A cell is blanked entirely (never partially rewritten) when it contains one of the known
unpublishable phrases - "blank is preferable to fabricated or unfinished copy". Structured
non-text columns (verified/verified_at/etc.) and source/provenance columns are left untouched.
"""

import openpyxl

WORKBOOK_PATH = "data/legacy-migration-batch-20/DestinationFinderAI_Legacy_Production_Batch_20_Reader_Ready_Authoritative_v3.3.xlsx"

KEYS = [
    "ajijic-mexico", "boquete-panama", "chiang-mai-thailand", "cuenca-ecuador", "da-nang-vietnam",
    "florianopolis-brazil", "funchal-portugal", "george-town-malaysia", "hua-hin-thailand", "lucca-italy",
    "merida-mexico", "monopoli-italy", "montevideo-uruguay", "nafplio-greece", "nice-france",
    "palm-springs-california-united-states", "paphos-cyprus", "santander-spain",
    "savannah-georgia-united-states", "sibenik-croatia",
]

PLACEHOLDER_PHRASES = [
    "not yet verified from an authoritative source",
    "verified imagery pending",
    "address-specific review recommended",
    "more local detail coming soon",
    "specific items to confirm",
]

# Columns that are provenance/source fields, not reader-facing copy - never blanked even if they
# happen to contain one of the phrases (they wouldn't, but this keeps the rule explicit and safe).
PROVENANCE_COLUMNS = {"source_name", "source_url", "verified", "verified_at", "confidence", "record_key", "destination_key", "methodology_version"}

# Sheets that carry destination-scoped reader-facing text (skip catalog/reference/meta sheets).
SKIP_SHEETS = {
    "README", "CATEGORIES", "SCORING_DIMENSIONS", "STAY_MODES", "SCHEMA_INDEX", "PREMIUM_REQUIREMENTS",
    "CHANGELOG", "PILOT_STATUS", "IMPORT_CONTRACT", "WORKBOOK_METADATA", "IMPORT_MANIFEST",
    "DESTINATION_ALIASES", "VALIDATION_RULES", "DATA_DICTIONARY",
}


def contains_placeholder(value) -> bool:
    if not isinstance(value, str):
        return False
    lower = value.lower()
    return any(phrase in lower for phrase in PLACEHOLDER_PHRASES)


def main():
    wb = openpyxl.load_workbook(WORKBOOK_PATH)
    total_blanked = 0
    per_sheet = {}

    for sheet_name in wb.sheetnames:
        if sheet_name in SKIP_SHEETS:
            continue
        ws = wb[sheet_name]
        header = [c.value for c in ws[1]]
        if "destination_key" not in header:
            continue
        idx = {h: i for i, h in enumerate(header) if h}
        dk_col = idx["destination_key"]
        blanked_here = 0
        for row in ws.iter_rows(min_row=2):
            if dk_col >= len(row):
                continue
            dest_key = row[dk_col].value
            if dest_key not in KEYS:
                continue
            for col_name, col_i in idx.items():
                if col_name in PROVENANCE_COLUMNS:
                    continue
                if col_i >= len(row):
                    continue
                cell = row[col_i]
                if contains_placeholder(cell.value):
                    cell.value = None
                    blanked_here += 1
        if blanked_here:
            per_sheet[sheet_name] = blanked_here
            total_blanked += blanked_here

    wb.save(WORKBOOK_PATH)
    print(f"Blanked {total_blanked} unpublishable cells across {len(per_sheet)} sheets:")
    for sheet_name, count in sorted(per_sheet.items(), key=lambda item: -item[1]):
        print(f"  {sheet_name}: {count}")


if __name__ == "__main__":
    main()
