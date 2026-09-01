"""
Bounded Batch 20 media completion: adds exactly 3 verified Wikimedia Commons gallery images per
destination (gallery_order 2, 3, 4; primary_image="NO"), bringing every destination from 1 image
(hero only) to the required 4 images (1 hero + 3 gallery). Candidates were discovered via the
Wikimedia Commons category/search API and license-verified via the imageinfo API
(scripts/commons_gallery_helper.py, scripts/find_batch20_gallery_images.py); the confirmed
candidate list lives in scripts/batch20_gallery_candidates.json.
"""

import json
import urllib.parse

import openpyxl
from datetime import date

WORKBOOK_PATH = "data/legacy-migration-batch-20/DestinationFinderAI_Legacy_Production_Batch_20_Reader_Ready_Authoritative_v3.3.xlsx"
TODAY = date.today().isoformat()


def special_filepath_url(bare_filename: str) -> str:
    return f"https://commons.wikimedia.org/wiki/Special:FilePath/{urllib.parse.quote(bare_filename)}"


def commons_file_url(bare_filename: str) -> str:
    return f"https://commons.wikimedia.org/wiki/File:{urllib.parse.quote(bare_filename)}"


def main():
    with open("scripts/batch20_gallery_candidates.json", encoding="utf-8") as f:
        candidates = json.load(f)

    wb = openpyxl.load_workbook(WORKBOOK_PATH)
    dest_ws = wb["DESTINATIONS"]
    dest_header = [c.value for c in dest_ws[1]]
    dest_idx = {h: i for i, h in enumerate(dest_header) if h}
    dest_names = {}
    for row in dest_ws.iter_rows(min_row=2):
        key = row[dest_idx["destination_key"]].value
        if key:
            dest_names[key] = row[dest_idx["destination_name"]].value

    media_ws = wb["MEDIA"]
    media_header = [c.value for c in media_ws[1]]
    media_idx = {h: i for i, h in enumerate(media_header) if h}
    next_row = media_ws.max_row + 1

    total_written = 0
    for dest_key, picks in candidates.items():
        if len(picks) != 3:
            print(f"SKIPPING {dest_key}: expected 3 candidates, found {len(picks)}")
            continue
        city_name = dest_names.get(dest_key, dest_key)
        for order_offset, pick in enumerate(picks, start=2):
            bare = pick["bare"]
            row_values = [None] * len(media_header)
            row_values[media_idx["media_key"]] = f"{dest_key}-gallery-{order_offset - 1:02d}"
            row_values[media_idx["destination_key"]] = dest_key
            row_values[media_idx["media_type"]] = "photo"
            row_values[media_idx["image_url"]] = special_filepath_url(bare)
            row_values[media_idx["caption"]] = city_name
            row_values[media_idx["subject"]] = city_name
            row_values[media_idx["primary_image"]] = "NO"
            row_values[media_idx["gallery_order"]] = order_offset
            row_values[media_idx["license_notes"]] = f"{pick['license_short']}, via Wikimedia Commons"
            row_values[media_idx["source_name"]] = "Wikimedia Commons"
            row_values[media_idx["source_url"]] = pick.get("url") or commons_file_url(bare)
            row_values[media_idx["verified"]] = True
            row_values[media_idx["verified_at"]] = TODAY
            row_values[media_idx["confidence"]] = "HIGH"
            for col_offset, value in enumerate(row_values):
                if value is not None:
                    media_ws.cell(row=next_row, column=col_offset + 1, value=value)
            next_row += 1
            total_written += 1

    wb.save(WORKBOOK_PATH)
    print(f"Wrote {total_written} new gallery MEDIA rows across {len(candidates)} destinations.")


if __name__ == "__main__":
    main()
