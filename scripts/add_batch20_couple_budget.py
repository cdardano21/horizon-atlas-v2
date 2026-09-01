"""
Bounded Batch 20 budget-profile addition: adds one new "couple" total_monthly_budget row per
destination, derived deterministically from each destination's EXISTING one-adult category ranges
(no new city-by-city cost research) using the task-specified multipliers:
housing x1.0, utilities x1.25, groceries x1.6, transportation x1.5, dining x1.6, healthcare x2.0,
leisure x1.6. Reuses the existing household_type column/convention already established for
COST_OF_LIVING rows (single/couple/family4) rather than inventing a second system.
"""

import math

import openpyxl
from datetime import date

WORKBOOK_PATH = "data/legacy-migration-batch-20/DestinationFinderAI_Legacy_Production_Batch_20_Reader_Ready_Authoritative_v3.3.xlsx"
TODAY = date.today().isoformat()

MULTIPLIERS = {
    "housing": 1.0,
    "utilities": 1.25,
    "groceries": 1.6,
    "transportation": 1.5,
    "dining": 1.6,
    "healthcare": 2.0,
    "leisure": 1.6,
}


def sensible_round(value: float) -> int:
    if value >= 1_000_000:
        unit = 50_000
    elif value >= 100_000:
        unit = 5_000
    elif value >= 10_000:
        unit = 500
    elif value >= 1_000:
        unit = 50
    else:
        unit = 5
    return int(round(value / unit) * unit)


def main():
    wb = openpyxl.load_workbook(WORKBOOK_PATH)
    ws = wb["COST_OF_LIVING"]
    header = [c.value for c in ws[1]]
    idx = {h: i for i, h in enumerate(header) if h}

    by_dest = {}
    for row in ws.iter_rows(min_row=2):
        dest_key = row[idx["destination_key"]].value
        if not dest_key:
            continue
        by_dest.setdefault(dest_key, []).append(row)

    next_row = ws.max_row + 1
    written = 0
    for dest_key, rows in by_dest.items():
        single_categories = {}
        total_row = None
        for row in rows:
            household = (row[idx["household_type"]].value or "").lower()
            category = (row[idx["category"]].value or "").lower()
            if household != "single":
                continue
            if category == "total_monthly_budget":
                total_row = row
            elif category in MULTIPLIERS:
                single_categories[category] = (row[idx["monthly_low"]].value, row[idx["monthly_high"]].value)

        if not total_row or len(single_categories) != len(MULTIPLIERS):
            print(f"SKIP {dest_key}: missing total row or category coverage ({len(single_categories)}/{len(MULTIPLIERS)})")
            continue

        couple_low_sum = sum(single_categories[cat][0] * mult for cat, mult in MULTIPLIERS.items())
        couple_high_sum = sum(single_categories[cat][1] * mult for cat, mult in MULTIPLIERS.items())
        couple_low = sensible_round(couple_low_sum)
        couple_high = sensible_round(couple_high_sum)

        currency = total_row[idx["currency"]].value
        source_name = total_row[idx["source_name"]].value
        source_url = total_row[idx["source_url"]].value

        row_values = [None] * len(header)
        row_values[idx["record_key"]] = f"{dest_key}-col-total_monthly_budget-couple"
        row_values[idx["destination_key"]] = dest_key
        row_values[idx["household_type"]] = "couple"
        row_values[idx["lifestyle_tier"]] = "comfortable"
        row_values[idx["category"]] = "total_monthly_budget"
        row_values[idx["monthly_low"]] = couple_low
        row_values[idx["monthly_high"]] = couple_high
        row_values[idx["currency"]] = currency
        row_values[idx["included_notes"]] = "Estimated for two adults sharing one home. Actual costs vary by housing, lifestyle and healthcare needs."
        row_values[idx["stay_mode_key"]] = "LONG_TERM_PERMANENT"
        row_values[idx["source_name"]] = source_name
        row_values[idx["source_url"]] = source_url
        row_values[idx["verified"]] = True
        row_values[idx["verified_at"]] = TODAY
        for col_offset, value in enumerate(row_values):
            if value is not None:
                ws.cell(row=next_row, column=col_offset + 1, value=value)
        next_row += 1
        written += 1

    wb.save(WORKBOOK_PATH)
    print(f"Wrote {written} couple total_monthly_budget rows.")


if __name__ == "__main__":
    main()
